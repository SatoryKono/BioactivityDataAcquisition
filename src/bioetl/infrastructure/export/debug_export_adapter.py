"""Infrastructure adapter for persisted debug export audit packs."""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import TYPE_CHECKING
from uuid import UUID

from bioetl.domain.types import DebugExportResult, RunID
from bioetl.infrastructure.control_plane import FileLineageStore
from bioetl.infrastructure.storage.atomic import atomic_write_text

if TYPE_CHECKING:
    from bioetl.domain.types import DebugExportPack

__all__ = ["DebugExportAdapter"]


class DebugExportAdapter:
    """Persist deterministic CSV/XLSX audit packs for debug export runs."""

    def __init__(
        self,
        *,
        lineage_store_path: str | None = None,
    ) -> None:
        self._lineage_store_path = lineage_store_path

    def write_pack(
        self,
        *,
        pack: DebugExportPack,
    ) -> DebugExportResult:
        root_path = self._resolve_root(pack)
        root_path.mkdir(parents=True, exist_ok=True)
        lineage_rows = tuple(self._load_lineage_rows(pack))
        tables = dict(pack.tables)
        if lineage_rows:
            tables["lineage"] = lineage_rows

        file_paths: list[str] = []
        canonical_artifacts: list[dict[str, object]] = []
        for table_name, rows in tables.items():
            csv_path = root_path / f"{table_name}.csv"
            schema_path = root_path / f"{table_name}.schema.json"
            self._write_csv(csv_path, rows, include_bom=pack.include_bom)
            self._write_schema(schema_path, rows)
            file_paths.extend((str(csv_path), str(schema_path)))
            canonical_artifacts.append(self._fingerprint(csv_path, root_path=root_path))
            canonical_artifacts.append(
                self._fingerprint(schema_path, root_path=root_path)
            )

        workbook_path: Path | None = None
        xlsx_skip_reason: str | None = None
        if "xlsx" in pack.formats:
            candidate_path = root_path / "debug_export.xlsx"
            try:
                self._write_xlsx(
                    candidate_path,
                    tables,
                    max_rows_per_sheet=pack.max_rows_per_sheet,
                )
            except ModuleNotFoundError as exc:
                if exc.name != "openpyxl":
                    raise
                xlsx_skip_reason = "openpyxl is not installed; CSV artifacts remain the source of truth"
            else:
                workbook_path = candidate_path
                file_paths.append(str(workbook_path))

        debug_export_hash = self._compute_pack_hash(canonical_artifacts)
        manifest_path = root_path / "manifest.json"
        manifest_payload = {
            "run_id": pack.run_id,
            "workflow_id": pack.workflow_id,
            "pipeline_id": pack.pipeline_id,
            "provider_id": pack.provider_id,
            "manifest_id": pack.manifest_id,
            "status": pack.status,
            "created_at": pack.created_at.isoformat(),
            "debug_export_hash": debug_export_hash,
            "xlsx_skip_reason": xlsx_skip_reason,
            "files": canonical_artifacts
            + (
                [
                    self._fingerprint(
                        workbook_path,
                        include_content_hash=False,
                        root_path=root_path,
                    )
                ]
                if workbook_path is not None
                else []
            ),
        }
        atomic_write_text(
            manifest_path,
            json.dumps(manifest_payload, indent=2, sort_keys=True) + "\n",
        )
        return DebugExportResult(
            root_path=str(root_path),
            manifest_path=str(manifest_path),
            debug_export_hash=debug_export_hash,
            file_paths=tuple(file_paths),
        )

    def _resolve_root(self, pack: DebugExportPack) -> Path:
        configured = Path(pack.output_root)
        if not configured.is_absolute():
            configured = Path.cwd() / configured
        return configured / pack.workflow_id / pack.pipeline_id / pack.run_id

    def _write_csv(
        self,
        output_path: Path,
        rows: tuple[dict[str, object], ...],
        *,
        include_bom: bool,
    ) -> None:
        headers = self._collect_headers(rows)
        encoding = "utf-8-sig" if include_bom else "utf-8"
        with output_path.open("w", encoding=encoding, newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=headers,
                extrasaction="ignore",
                quoting=csv.QUOTE_ALL,
            )
            writer.writeheader()
            for row in rows:
                writer.writerow(
                    {
                        header: self._normalize_csv_value(row.get(header))
                        for header in headers
                    }
                )

    def _write_schema(
        self,
        output_path: Path,
        rows: tuple[dict[str, object], ...],
    ) -> None:
        headers = self._collect_headers(rows)
        types = {
            header: sorted(
                {
                    type(row.get(header)).__name__
                    for row in rows
                    if row.get(header) is not None
                }
            )
            for header in headers
        }
        atomic_write_text(
            output_path,
            json.dumps(
                {"columns": [{"name": h, "types": types[h]} for h in headers]},
                indent=2,
                sort_keys=True,
            )
            + "\n",
        )

    def _write_xlsx(
        self,
        output_path: Path,
        tables: dict[str, tuple[dict[str, object], ...]],
        *,
        max_rows_per_sheet: int,
    ) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.properties.created = None
        workbook.properties.modified = None
        first_sheet = True
        for table_name, rows in tables.items():
            headers = self._collect_headers(rows)
            chunk_size = max(1, min(max_rows_per_sheet - 1, 1_000_000))
            chunks = [rows[i : i + chunk_size] for i in range(0, len(rows), chunk_size)]
            if not chunks:
                chunks = [()]
            for chunk_index, chunk in enumerate(chunks, start=1):
                sheet_name = (
                    table_name
                    if len(chunks) == 1
                    else f"{table_name}_{chunk_index:04d}"
                )[:31]
                if first_sheet:
                    worksheet = workbook.active
                    worksheet.title = sheet_name
                    first_sheet = False
                else:
                    worksheet = workbook.create_sheet(title=sheet_name)
                worksheet.freeze_panes = "A2"
                worksheet.append(headers)
                for row in chunk:
                    worksheet.append(
                        [
                            self._normalize_csv_value(row.get(header))
                            for header in headers
                        ]
                    )
                worksheet.auto_filter.ref = worksheet.dimensions
        workbook.save(output_path)

    def _collect_headers(
        self,
        rows: tuple[dict[str, object], ...],
    ) -> list[str]:
        headers: list[str] = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
        return headers

    def _normalize_csv_value(self, value: object | None) -> object:
        if value is None:
            return ""
        if isinstance(value, (dict, list, tuple)):
            return json.dumps(value, ensure_ascii=False, sort_keys=True)
        return value

    def _fingerprint(
        self,
        path: Path,
        *,
        include_content_hash: bool = True,
        root_path: Path | None = None,
    ) -> dict[str, object]:
        payload = path.read_bytes()
        relative_path = (
            str(path.relative_to(root_path)) if root_path is not None else str(path)
        )
        result = {
            "path": relative_path,
            "size_bytes": len(payload),
        }
        if include_content_hash:
            result["sha256"] = hashlib.sha256(payload).hexdigest()
        return result

    def _compute_pack_hash(
        self,
        artifacts: list[dict[str, object]],
    ) -> str:
        payload = json.dumps(artifacts, sort_keys=True, separators=(",", ":")).encode(
            "utf-8"
        )
        return hashlib.sha256(payload).hexdigest()

    def _load_lineage_rows(self, pack: DebugExportPack) -> list[dict[str, object]]:
        if self._lineage_store_path is None:
            return []
        try:
            store = FileLineageStore(base_path=Path(self._lineage_store_path))
            fragments = (
                store.list_by_manifest_id(pack.manifest_id)
                if pack.manifest_id is not None
                else store.list_by_run_id(RunID(UUID(pack.run_id)))
            )
        except (OSError, TypeError, ValueError):
            return []
        rows: list[dict[str, object]] = []
        for fragment in fragments:
            for node in fragment.nodes:
                rows.append(
                    {
                        "fragment_id": fragment.fragment_id,
                        "stored_fragment_id": fragment.stored_fragment_id,
                        "manifest_id": fragment.manifest_id,
                        "run_id": fragment.run_id,
                        "node_id": node.node_id,
                        "edge_type": "",
                        "related_node_id": "",
                        "node_type": node.node_type,
                    }
                )
            for edge in fragment.edges:
                rows.append(
                    {
                        "fragment_id": fragment.fragment_id,
                        "stored_fragment_id": fragment.stored_fragment_id,
                        "manifest_id": fragment.manifest_id,
                        "run_id": fragment.run_id,
                        "node_id": edge.source.node_id,
                        "edge_type": edge.edge_type,
                        "related_node_id": edge.target.node_id,
                        "node_type": "",
                    }
                )
        return rows
